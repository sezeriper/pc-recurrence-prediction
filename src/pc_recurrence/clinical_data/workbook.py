from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pc_recurrence.image_data.workbook import EXPECTED_HEADERS


@dataclass(frozen=True)
class ClinicalWorkbookRow:
    patient_id: str
    row_number: int
    hospital_number: str | float | None
    surgery: Any
    pathology: Any
    recurrence_raw: Any
    age: Any
    ca_19_9: Any
    total_bilirubin: Any
    direct_bilirubin: Any
    symptom: Any
    serum_albumin: Any
    absolute_lymphocyte: Any
    crp: Any
    cally: Any
    pni: Any
    image_range: Any
    ct_report: Any

    @property
    def dicom_folder(self) -> str | None:
        value = _clean_scalar(self.hospital_number)
        if value is None:
            return None
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]
        return f"PATIENT{text}" if text else None


def _clean_scalar(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def load_clinical_workbook(
    workbook_path: str | Path, sheet_name: str = "Sayfa1"
) -> list[ClinicalWorkbookRow]:
    """Load the complete clinical schema without changing the source workbook."""
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Expected worksheet {sheet_name!r}; found {workbook.sheetnames}")
        sheet = workbook[sheet_name]
        raw_rows = sheet.iter_rows(min_col=1, max_col=len(EXPECTED_HEADERS), values_only=True)
        raw_headers = next(raw_rows, None)
        if raw_headers is None:
            raise ValueError("Workbook is empty")
        headers = tuple(str(value).strip() if value is not None else "" for value in raw_headers)
        if headers != EXPECTED_HEADERS:
            raise ValueError(
                f"Unexpected workbook schema. Expected {EXPECTED_HEADERS}; found {headers}"
            )

        records: list[ClinicalWorkbookRow] = []
        seen_ids: set[str] = set()
        for row_number, raw_row in enumerate(raw_rows, start=2):
            cleaned = tuple(_clean_scalar(value) for value in raw_row)
            if all(value is None for value in cleaned):
                continue
            patient_id = cleaned[0]
            if not isinstance(patient_id, str) or not patient_id:
                raise ValueError(f"Row {row_number} has no valid patient ID")
            if patient_id in seen_ids:
                raise ValueError(f"Duplicate patient ID {patient_id!r} at row {row_number}")
            seen_ids.add(patient_id)
            records.append(
                ClinicalWorkbookRow(
                    patient_id=patient_id,
                    row_number=row_number,
                    hospital_number=cleaned[1],
                    surgery=cleaned[2],
                    pathology=cleaned[3],
                    recurrence_raw=cleaned[4],
                    age=cleaned[5],
                    ca_19_9=cleaned[6],
                    total_bilirubin=cleaned[7],
                    direct_bilirubin=cleaned[8],
                    symptom=cleaned[9],
                    serum_albumin=cleaned[10],
                    absolute_lymphocyte=cleaned[11],
                    crp=cleaned[12],
                    cally=cleaned[13],
                    pni=cleaned[14],
                    image_range=cleaned[15],
                    ct_report=cleaned[16],
                )
            )
    finally:
        workbook.close()

    if not records:
        raise ValueError("No populated patient rows were found")
    return records


def select_clinical_workbook_rows(
    rows: list[ClinicalWorkbookRow], patients: set[str] | None
) -> list[ClinicalWorkbookRow]:
    """Select rows by workbook patient ID or anonymized DICOM-folder alias."""
    if patients is None:
        return rows
    selected = [
        row
        for row in rows
        if row.patient_id in patients
        or (row.dicom_folder is not None and row.dicom_folder in patients)
    ]
    matched = {row.patient_id for row in selected}
    matched.update(row.dicom_folder for row in selected if row.dicom_folder is not None)
    unknown = sorted(patients - matched)
    if unknown:
        raise ValueError(f"Unknown patient aliases: {', '.join(unknown)}")
    return selected
